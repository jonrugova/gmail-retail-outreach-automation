import base64
import csv
import os
import re
import time
from email.message import EmailMessage
from email.utils import parseaddr
from pathlib import Path

from openpyxl import load_workbook

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError


# =========================================================
# CONFIGURATION
# =========================================================

EXCEL_FILE = "retail_outreach_tracker.xlsx"
SHEET_NAME = "Retail Outreach"

# True = finds threads but sends NOTHING.
# False = actually sends emails.
DRY_RUN = True

# Limits how many unsent contacts are processed per run,
# including dry runs. This prevents Gmail API quota spikes.
MAX_CONTACTS_PER_RUN = 30

# Once sending is enabled, never send more than this per run.
MAX_SEND_PER_RUN = 30

# Keep Gmail result sets small.
MAX_SEARCH_RESULTS_PER_QUERY = 10
MAX_CANDIDATES_PER_CONTACT = 15

# Small pauses reduce burst traffic.
API_CALL_DELAY = 0.8
DELAY_BETWEEN_CONTACTS = 2

# Backoff when Gmail temporarily rate-limits requests.
RATE_LIMIT_RETRY_DELAYS = [10, 20, 40]

SCOPES = [
    "https://www.googleapis.com/auth/gmail.modify"
]

OUR_EMAILS = {
    "team@example.com",
}

MARKER_TEXT = (
    "Following recent changes within our retail partnerships team"
)

KEYWORDS = [
    "wholesale",
    "retailer",
    "retail partnership",
    "stockist",
    "reseller",
    "distributor",
    "carry your",
    "carry the brand",
    "carry your brand",
    "carry your collection",
    "wholesale pricing",
    "wholesale price",
    "wholesale order",
    "become a retailer",
    "becoming a retailer",
]

GMAIL_RETAIL_FILTER = (
    "{wholesale retailer retail stockist reseller distributor "
    "partnership \"carry your\"}"
)

TEMPLATE = """{greeting}

Thank you for previously reaching out to us regarding a potential retail partnership.

Following recent changes within our retail partnerships team, we are reconnecting with retailers and distributors who previously expressed interest in carrying our collections.

We are currently extending special pricing to selected retail partners. If you are still interested in working with us, we would be happy to share our latest lookbook, retail partner pricing, and current partnership terms.

We would be delighted to hear from you.

Warm regards,
Retail Partnerships Team
"""


# =========================================================
# AUTHENTICATION
# =========================================================

def get_gmail_service():
    creds = None

    if os.path.exists("token.json"):
        creds = Credentials.from_authorized_user_file(
            "token.json", SCOPES
        )

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(
                "credentials.json",
                SCOPES,
            )
            creds = flow.run_local_server(port=0)

        with open("token.json", "w") as token:
            token.write(creds.to_json())

    return build("gmail", "v1", credentials=creds)


# =========================================================
# RATE-LIMIT SAFE GMAIL EXECUTION
# =========================================================

def is_rate_limit_error(error):
    status = getattr(getattr(error, "resp", None), "status", None)
    text = str(error).lower()

    return (
        status in (403, 429)
        and (
            "ratelimitexceeded" in text
            or "userratelimitexceeded" in text
            or "quota exceeded" in text
            or "rate limit" in text
        )
    )


def gmail_execute(request, label="Gmail API request"):
    attempts = 1 + len(RATE_LIMIT_RETRY_DELAYS)

    for attempt in range(attempts):
        try:
            result = request.execute()
            time.sleep(API_CALL_DELAY)
            return result

        except HttpError as error:
            if (
                is_rate_limit_error(error)
                and attempt < len(RATE_LIMIT_RETRY_DELAYS)
            ):
                wait_seconds = RATE_LIMIT_RETRY_DELAYS[attempt]
                print(
                    f"Rate limit reached during {label}. "
                    f"Waiting {wait_seconds}s, then retrying..."
                )
                time.sleep(wait_seconds)
                continue

            raise


# =========================================================
# GMAIL HELPERS
# =========================================================

def get_headers(message):
    headers = {}

    for header in message.get("payload", {}).get("headers", []):
        headers[header["name"].lower()] = header["value"]

    return headers


def decode_body_data(data):
    if not data:
        return ""

    try:
        return base64.urlsafe_b64decode(
            data + "=" * (-len(data) % 4)
        ).decode("utf-8", errors="replace")
    except Exception:
        return ""


def extract_body_from_payload(payload):
    mime_type = payload.get("mimeType", "")

    body_data = payload.get("body", {}).get("data")
    if body_data:
        text = decode_body_data(body_data)

        if mime_type == "text/html":
            text = re.sub(r"<[^>]+>", " ", text)

        return text

    texts = []

    for part in payload.get("parts", []):
        text = extract_body_from_payload(part)

        if text:
            texts.append(text)

    return "\n".join(texts)


def get_message(service, message_id):
    request = (
        service.users()
        .messages()
        .get(
            userId="me",
            id=message_id,
            format="full",
        )
    )

    return gmail_execute(
        request,
        label=f"message.get {message_id}",
    )


def sender_email(message):
    headers = get_headers(message)

    return parseaddr(
        headers.get("from", "")
    )[1].lower()


def message_subject(message):
    return get_headers(message).get("subject", "")


# =========================================================
# MATCHING THE CORRECT RETAIL THREAD
# =========================================================

def score_message(message, customer_email):
    headers = get_headers(message)

    sender = parseaddr(
        headers.get("from", "")
    )[1].lower()

    subject = headers.get("subject", "")
    body = extract_body_from_payload(
        message.get("payload", {})
    )

    text = f"{subject}\n{body}".lower()

    score = 0

    if sender == customer_email.lower():
        score += 100

    if customer_email.lower() in text:
        score += 40

    if "SENT" in message.get("labelIds", []):
        score -= 100

    for keyword in KEYWORDS:
        if keyword in text:
            score += 12

    subject_lower = subject.lower()

    if any(
        word in subject_lower
        for word in [
            "wholesale",
            "retailer",
            "retail",
            "stockist",
            "reseller",
            "partnership",
        ]
    ):
        score += 25

    return score


def list_message_ids(service, query):
    request = (
        service.users()
        .messages()
        .list(
            userId="me",
            q=query,
            maxResults=MAX_SEARCH_RESULTS_PER_QUERY,
        )
    )

    result = gmail_execute(
        request,
        label=f"messages.list [{query}]",
    )

    return [
        item["id"]
        for item in result.get("messages", [])
    ]


def search_candidate_messages(service, email):
    ordered_ids = []
    seen = set()

    targeted_queries = [
        f'from:{email} {GMAIL_RETAIL_FILTER} -in:spam -in:trash',
        f'"{email}" {GMAIL_RETAIL_FILTER} -in:spam -in:trash',
    ]

    for query in targeted_queries:
        for message_id in list_message_ids(service, query):
            if message_id not in seen:
                seen.add(message_id)
                ordered_ids.append(message_id)

            if len(ordered_ids) >= MAX_CANDIDATES_PER_CONTACT:
                return ordered_ids

    if not ordered_ids:
        fallback_queries = [
            f'from:{email} -in:spam -in:trash',
            f'"{email}" -in:spam -in:trash',
        ]

        for query in fallback_queries:
            for message_id in list_message_ids(service, query):
                if message_id not in seen:
                    seen.add(message_id)
                    ordered_ids.append(message_id)

                if len(ordered_ids) >= MAX_CANDIDATES_PER_CONTACT:
                    return ordered_ids

    return ordered_ids


def find_best_thread(service, customer_email):
    candidate_ids = search_candidate_messages(
        service,
        customer_email,
    )

    best_message = None
    best_score = -999999

    for message_id in candidate_ids:
        message = get_message(
            service,
            message_id,
        )

        score = score_message(
            message,
            customer_email,
        )

        if score > best_score:
            best_score = score
            best_message = message

    if not best_message or best_score < 20:
        return None, best_score

    thread_id = best_message["threadId"]

    request = (
        service.users()
        .threads()
        .get(
            userId="me",
            id=thread_id,
            format="full",
        )
    )

    thread = gmail_execute(
        request,
        label=f"threads.get {thread_id}",
    )

    return thread, best_score


# =========================================================
# DUPLICATE PROTECTION
# =========================================================

def campaign_already_sent(thread):
    for message in thread.get("messages", []):
        if "SENT" not in message.get("labelIds", []):
            continue

        body = extract_body_from_payload(
            message.get("payload", {})
        )

        if MARKER_TEXT in body:
            return True

    return False


# =========================================================
# FIND BEST MESSAGE TO REPLY TO
# =========================================================

def latest_inbound_message(thread):
    inbound = []

    for message in thread.get("messages", []):
        headers = get_headers(message)

        sender = parseaddr(
            headers.get("from", "")
        )[1].lower()

        is_sent = "SENT" in message.get(
            "labelIds",
            [],
        )

        if not is_sent and sender not in OUR_EMAILS:
            inbound.append(message)

    if not inbound:
        return None

    inbound.sort(
        key=lambda x: int(x.get("internalDate", 0))
    )

    return inbound[-1]


# =========================================================
# PERSONALIZATION
# =========================================================

def determine_greeting(message):
    body = extract_body_from_payload(
        message.get("payload", {})
    )

    match = re.search(
        r"Name:\s*\n+\s*([^\n]+)",
        body,
        flags=re.IGNORECASE,
    )

    if match:
        full_name = match.group(1).strip()
        first = full_name.split()[0]

        if first:
            return f"Dear {first},"

    match = re.search(
        r"\bmy name is\s+([A-Za-zÀ-ÿ'’-]+)",
        body,
        flags=re.IGNORECASE,
    )

    if match:
        return f"Dear {match.group(1)},"

    return "Hello,"


# =========================================================
# SEND THREADED REPLY
# =========================================================

def send_reply(
    service,
    recipient,
    original_message,
):
    headers = get_headers(original_message)

    subject = headers.get("subject", "")
    message_id_header = headers.get("message-id", "")
    references = headers.get("references", "")

    if not subject:
        raise ValueError("Original message has no Subject header.")

    if not message_id_header:
        raise ValueError(
            "Original message has no Message-ID header."
        )

    greeting = determine_greeting(
        original_message
    )

    text = TEMPLATE.format(
        greeting=greeting
    )

    msg = EmailMessage()

    msg["To"] = recipient
    msg["Subject"] = subject
    msg["In-Reply-To"] = message_id_header

    if references:
        msg["References"] = (
            references + " " + message_id_header
        )
    else:
        msg["References"] = message_id_header

    msg.set_content(text)

    encoded_message = base64.urlsafe_b64encode(
        msg.as_bytes()
    ).decode()

    api_message = {
        "raw": encoded_message,
        "threadId": original_message["threadId"],
    }

    request = (
        service.users()
        .messages()
        .send(
            userId="me",
            body=api_message,
        )
    )

    return gmail_execute(
        request,
        label=f"messages.send to {recipient}",
    )


# =========================================================
# LOGGING
# =========================================================

def log_result(
    email,
    status,
    thread_id="",
    subject="",
    gmail_message_id="",
    detail="",
):
    log_exists = Path(
        "retail_reply_log.csv"
    ).exists()

    with open(
        "retail_reply_log.csv",
        "a",
        newline="",
        encoding="utf-8",
    ) as f:
        writer = csv.writer(f)

        if not log_exists:
            writer.writerow([
                "Email",
                "Status",
                "Thread ID",
                "Subject",
                "Gmail Message ID",
                "Detail",
            ])

        writer.writerow([
            email,
            status,
            thread_id,
            subject,
            gmail_message_id,
            detail,
        ])


# =========================================================
# MAIN
# =========================================================

def main():
    service = get_gmail_service()

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook[SHEET_NAME]

    headers = {}

    for cell in sheet[1]:
        if cell.value:
            headers[str(cell.value).strip()] = cell.column

    email_col = headers["Email"]
    sent_col = headers["Email Sent"]

    processed_this_run = 0
    sent_this_run = 0

    for row in range(2, sheet.max_row + 1):
        email = sheet.cell(
            row,
            email_col,
        ).value

        status = sheet.cell(
            row,
            sent_col,
        ).value

        if not email:
            continue

        email = str(email).strip().lower()

        if str(status).strip().lower() == "yes":
            print(f"SKIP already sent: {email}")
            continue

        if (
            MAX_CONTACTS_PER_RUN
            and processed_this_run >= MAX_CONTACTS_PER_RUN
        ):
            print(
                f"Reached MAX_CONTACTS_PER_RUN="
                f"{MAX_CONTACTS_PER_RUN}."
            )
            break

        processed_this_run += 1

        print()
        print("=" * 70)
        print(
            f"[{processed_this_run}/{MAX_CONTACTS_PER_RUN}] "
            f"{email}"
        )

        try:
            thread, score = find_best_thread(
                service,
                email,
            )

            if not thread:
                print(
                    f"NO MATCH (score {score})"
                )

                log_result(
                    email,
                    "NO MATCH",
                    detail=f"score={score}",
                )

                time.sleep(DELAY_BETWEEN_CONTACTS)
                continue

            if campaign_already_sent(thread):
                print(
                    "Campaign already exists in this thread."
                )

                log_result(
                    email,
                    "ALREADY SENT",
                    thread["id"],
                    detail=f"score={score}",
                )

                if not DRY_RUN:
                    sheet.cell(
                        row,
                        sent_col,
                    ).value = "Yes"

                    workbook.save(
                        EXCEL_FILE
                    )

                time.sleep(DELAY_BETWEEN_CONTACTS)
                continue

            original = latest_inbound_message(
                thread
            )

            if not original:
                print("No inbound message found.")

                log_result(
                    email,
                    "NO INBOUND",
                    thread["id"],
                    detail=f"score={score}",
                )

                time.sleep(DELAY_BETWEEN_CONTACTS)
                continue

            subject = message_subject(original)

            print(
                f"Thread: {thread['id']}"
            )

            print(
                f"Subject: {subject}"
            )

            print(
                f"Match score: {score}"
            )

            if DRY_RUN:
                print(
                    "DRY RUN — email NOT sent."
                )

                log_result(
                    email,
                    "DRY RUN",
                    thread["id"],
                    subject,
                    detail=f"score={score}",
                )

                time.sleep(DELAY_BETWEEN_CONTACTS)
                continue

            result = send_reply(
                service,
                email,
                original,
            )

            gmail_message_id = result["id"]

            sheet.cell(
                row,
                sent_col,
            ).value = "Yes"

            workbook.save(
                EXCEL_FILE
            )

            log_result(
                email,
                "SENT",
                thread["id"],
                subject,
                gmail_message_id,
            )

            print(
                f"SENT → {gmail_message_id}"
            )

            sent_this_run += 1

            if (
                MAX_SEND_PER_RUN
                and sent_this_run >= MAX_SEND_PER_RUN
            ):
                print(
                    "Reached MAX_SEND_PER_RUN."
                )
                break

            time.sleep(DELAY_BETWEEN_CONTACTS)

        except HttpError as error:
            print(
                f"GMAIL ERROR: {error}"
            )

            log_result(
                email,
                "ERROR",
                detail=str(error),
            )

            time.sleep(DELAY_BETWEEN_CONTACTS)

        except Exception as error:
            print(
                f"ERROR: {error}"
            )

            log_result(
                email,
                "ERROR",
                detail=str(error),
            )

            time.sleep(DELAY_BETWEEN_CONTACTS)


if __name__ == "__main__":
    main()
